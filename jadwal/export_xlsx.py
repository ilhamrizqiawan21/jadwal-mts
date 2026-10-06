"""Ekspor Excel (.xlsx): dua lembar, tata letak sama dengan cetakan PDF.

Ukuran (lebar kolom, tinggi baris, huruf) diskalakan SCALE kali supaya enak
dibaca di Excel; pengaturan cetak "muat satu halaman" mengembalikannya ke
ukuran kertas yang sama dengan PDF.
"""
from __future__ import annotations

import base64
import io
from typing import List

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from .export_pdf import (JUDUL_HARI, _E14, _P1, _X, _col_edges, _ket_range, _kode_font,
                         _total, HEAD_BLUE)
from .model import Proyek

SCALE = 1.6
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NOB = Border()


def _w(pt: float) -> float:
    """Lebar kolom Excel (karakter) dari titik, sudah diskalakan."""
    px = pt * SCALE / 0.75
    return max(0.5, (px - 5) / 7)


def _txt(v):
    """Teks aman untuk sel: buang karakter kontrol yang ditolak Excel."""
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def _fill(hexcol):
    if not hexcol:
        return None
    return PatternFill("solid", start_color=hexcol.lstrip("#").upper(), end_color=hexcol.lstrip("#").upper())


def _font(name, size, bold=True, color="000000"):
    return Font(name=name, size=round(size * SCALE * 2) / 2, bold=bold, color=color)


def _put(ws, r, c, v=None, font=None, fill=None, align="center", border=BOX, rot=None, wrap=False,
         valign="center", r2=None, c2=None, shrink=False):
    """Isi sel (atau rentang gabungan) dengan gaya."""
    r2, c2 = r2 or r, c2 or c
    if (r2, c2) != (r, c):
        ws.merge_cells(start_row=r, start_column=c, end_row=r2, end_column=c2)
    for rr in range(r, r2 + 1):
        for cc in range(c, c2 + 1):
            cell = ws.cell(rr, cc)
            if border is not None:
                cell.border = border
            if fill is not None:
                cell.fill = fill
    cell = ws.cell(r, c)
    if v is not None:
        cell.value = _txt(v)
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"        # teks operator tidak boleh dibaca sebagai formula
    if font:
        cell.font = font
    cell.alignment = Alignment(horizontal=align, vertical=valign, text_rotation=rot or 0, wrap_text=wrap,
                               shrink_to_fit=shrink)
    return cell


def _img(ws, b64, anchor, w_pt, h_pt):
    if not b64:
        return
    try:
        im = XLImage(io.BytesIO(base64.b64decode(b64)))
    except Exception:
        return
    im.width, im.height = w_pt * SCALE / 0.75, h_pt * SCALE / 0.75
    ws.add_image(im, anchor)


# ================================================================== BEBAN
def _sheet_beban(wb: Workbook, p: Proyek):
    ws = wb.active
    ws.title = "Beban Mengajar"
    g = _P1
    n = len(p.kelas)
    edges = _col_edges(g["x_cls0"], g["x_cls1"], n)
    # kolom: A NO | B nama | C kode | D bidang | E BBT | F.. kelas | wali | KET
    ws.column_dimensions["A"].width = _w(g["x_nama"] - g["x_no"])
    ws.column_dimensions["B"].width = _w(g["x_kode"] - g["x_nama"])
    ws.column_dimensions["C"].width = _w(g["x_bid"] - g["x_kode"])
    ws.column_dimensions["D"].width = _w(g["x_bbt"] - g["x_bid"])
    ws.column_dimensions["E"].width = _w(g["x_cls0"] - g["x_bbt"])
    for i in range(n):
        ws.column_dimensions[L(6 + i)].width = _w(edges[i + 1] - edges[i])
    cw = 6 + n
    ws.column_dimensions[L(cw)].width = _w(g["x_wali1"] - g["x_cls1"])
    ws.column_dimensions[L(cw + 1)].width = _w(g["x_ket1"] - g["x_wali1"])
    last = cw + 1
    m = p.meta

    # kop + judul
    ws.row_dimensions[1].height = 82 * SCALE
    _img(ws, m.kop_png, "A1", 523, 81)
    for r, (txt, size) in enumerate(((m.judul_beban, 14.4), (m.judul_beban2, 9.2),
                                     (f"SEMESTER {m.semester} TAHUN PELAJARAN {m.tahun}", 13.0)), start=2):
        ws.row_dimensions[r].height = 15 * SCALE
        _put(ws, r, 1, txt, _font("Tahoma", size), border=None, c2=last)
    ws.row_dimensions[5].height = 6 * SCALE
    # kepala tabel
    h0, h1, h2 = 6, 7, 8
    ws.row_dimensions[h0].height = 10.3 * SCALE
    ws.row_dimensions[h1].height = 74.7 * SCALE
    bold = _font("Tahoma", 9.2)
    _put(ws, h0, 1, "NO.", _font("Tahoma", 9.2, False), r2=h1, align="left")
    _put(ws, h0, 2, "NAMA GURU", bold, r2=h1, align="left")
    _put(ws, h0, 3, "KODE", bold, r2=h1)
    _put(ws, h0, 4, "BIDANG STUDI", bold, r2=h1)
    _put(ws, h0, 5, "BBT", bold)
    _put(ws, h0, 6, "KELAS", bold, c2=5 + n)
    _put(ws, h1, 5, "JAM", _font("Tahoma", 7.8))
    for i, k in enumerate(p.kelas):
        _put(ws, h1, 6 + i, k.nama, bold, fill=_fill("#DCE6F0"), rot=90)
    _put(ws, h0, cw, "WALI KELAS", bold, r2=h1, rot=90)
    _put(ws, h0, cw + 1, "KET", _font("Tahoma", 7.8), r2=h1)

    rows = list(p.penugasan)
    nrow = len(rows)
    r0 = h1 + 1
    # wali per baris
    wali_row = {}
    for k in p.kelas:
        if not k.wali:
            continue
        cand = [i for i, a in enumerate(rows) if a.guru == k.wali]
        if not cand:
            continue
        pick = next((i for i in cand if rows[i].beban.get(k.id, 0) > 0), cand[0])
        wali_row[pick] = (wali_row.get(pick, "") + " " + k.nama).strip()

    for i, a in enumerate(rows):
        r = r0 + i
        ws.row_dimensions[r].height = 14.4 * SCALE
        special = a.fill_baris
        _put(ws, r, 4, a.mapel, _font("Tahoma", a.ukuran_mapel or 7.8),
             fill=_fill(a.fill_mapel or special), wrap=True)
        fnt = _kode_font(a.kode)
        _put(ws, r, 3, a.kode, _font("Tahoma" if fnt.startswith("Tahoma") else "Arial",
                                     7.8 if a.kode in ("1", "2") else 9.2),
             fill=_fill(special if special else (a.warna if a.kode else None)))
        _put(ws, r, 5, None, _font("Tahoma", 9.2), fill=_fill(special))
        kd, ke = _ket_range(a, n)
        for j in range(n):
            if a.keterangan and kd <= j <= ke:
                continue
            v = a.beban.get(p.kelas[j].id)
            fill = _fill(special) if special else (PatternFill("solid", start_color="000000", end_color="000000")
                                                   if v is None else None)
            c = _put(ws, r, 6 + j, v if v else None, _font("Tahoma", 7.8, bool(a.merah),
                                                        "FF0000" if a.merah else "000000"), fill=fill)
        if a.keterangan:
            _put(ws, r, 6 + kd, a.keterangan.strip(), _font("Tahoma", 9.2), fill=_fill(special),
                 align="left", c2=6 + ke, shrink=True)
        w = wali_row.get(i, "")
        _put(ws, r, cw, w or None, _font("Tahoma", 11.8),
             fill=_fill(special) if special else (None if w else PatternFill("solid", start_color="BFBFBF", end_color="BFBFBF")))
        tot = a.ket_manual if a.ket_manual is not None else _total(a)
        _put(ws, r, cw + 1, tot or None, _font("Tahoma", 10.4, True, "FF0000" if a.ket_merah else "000000"),
             fill=_fill(special))

    # BBT (gabung bila diminta) -> tulis setelah baris
    for i, a in enumerate(rows):
        if a.bbt is None or (a.bbt_gabung and i > 0):
            continue
        r = r0 + i
        end = r + 1 if i + 1 < nrow and rows[i + 1].bbt_gabung else r
        if end != r:
            ws.merge_cells(start_row=r, start_column=5, end_row=end, end_column=5)
        c = ws.cell(r, 5)
        c.value = a.bbt
        c.font = _font("Tahoma", 9.2)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX

    # NO & nama (digabung per guru)
    i, no = 0, 0
    while i < nrow:
        j = i
        while j + 1 < nrow and rows[j + 1].guru == rows[i].guru:
            j += 1
        no += 1
        guru = p.guru_by_id(rows[i].guru)
        _put(ws, r0 + i, 1, no, _font("Tahoma", 7.8, False), r2=r0 + j,
             fill=_fill(guru.fill_no) if guru and guru.fill_no else None)
        # nama: baris ber-nama_tampil dicetak sendiri
        k = i
        while k <= j:
            a = rows[k]
            if a.nama_tampil:
                _put(ws, r0 + k, 2, a.nama_tampil, _font("Tahoma", 7.8, False), align="left",
                     fill=_fill(a.fill_baris))
                k += 1
            else:
                e = k
                while e + 1 <= j and not rows[e + 1].nama_tampil:
                    e += 1
                _put(ws, r0 + k, 2, guru.nama if guru else "", _font("Tahoma", 7.8, False), align="left",
                     r2=r0 + e, fill=_fill(guru.fill_nama) if guru and guru.fill_nama else _fill(rows[k].fill_baris),
                     wrap=True)
                k = e + 1
        i = j + 1

    # tanda tangan
    t = m.beban_ttd
    rs = r0 + nrow + 1
    jab = t.get("jabatan", "")
    ws.row_dimensions[rs].height = (28 if jab else 20) * SCALE
    tt = t.get("tempat_tanggal", "")
    c_tt = _put(ws, rs, cw - 6, tt, _font("Tahoma", 9.2, False), border=None, align="left", c2=last,
                wrap=bool(jab))
    if jab:     # jabatan tebal & lebih kecil pada baris kedua sel yang sama (seperti cetakan PDF)
        c_tt.value = CellRichText(
            TextBlock(InlineFont(rFont="Tahoma", sz=round(9.2 * SCALE * 2) / 2), _txt(tt) + "\n"),
            TextBlock(InlineFont(rFont="Tahoma", sz=round(7.0 * SCALE * 2) / 2, b=True), _txt(jab)))
    ws.row_dimensions[rs + 1].height = 64 * SCALE
    _img(ws, m.ttd_png, f"{L(cw - 6)}{rs + 1}", 89, 51)
    ws.row_dimensions[rs + 2].height = 11 * SCALE
    nm = _put(ws, rs + 2, cw - 6, t.get("nama", ""), _font("Calibri", 7.8), border=None, align="left", c2=last)
    nm.font = Font(name="Calibri", size=round(7.8 * SCALE * 2) / 2, bold=True, underline="single")
    _put(ws, rs + 3, cw - 6, t.get("nip", ""), _font("Calibri", 7.8), border=None, align="left", c2=last)
    ws.page_setup.paperSize = 5
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{L(last)}{rs + 3}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.3
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False


# ================================================================== JADWAL
def _expand(rows):
    out = []
    for r in rows:
        if r["t"] == "blok":
            out += [("jp", r, 0), ("jp", r, 1)]
        elif r["t"] == "tutup" and r.get("jam"):
            out += [("tutup_jp", r, k) for k in range(len(r["jam"]))]
        else:
            out.append((r["t"], r, 0))
    return out


def _sheet_jadwal(wb: Workbook, p: Proyek):
    ws = wb.create_sheet("Jadwal")
    n = len(p.kelas)
    by_kode = p.penugasan_by_kode()
    m = p.meta
    from .export_pdf import _edges_hari
    # kolom: A NO | B WAKTU | blok1 (n) | sep | blok2 (n) | blok3 (n)
    c1, c2, c3 = 3, 3 + n + 1, 3 + 2 * n + 1
    ws.column_dimensions["A"].width = _w(10.6)
    ws.column_dimensions["B"].width = _w(36.1)
    for start, hari in ((c1, "Selasa"), (c2, "Rabu"), (c3, "Kamis")):
        e = _edges_hari(hari, n)
        for i in range(n):
            ws.column_dimensions[L(start + i)].width = _w(e[i + 1] - e[i])
    ws.column_dimensions[L(c1 + n)].width = _w(4.4)
    last_col = c3 + n - 1
    blocks = {"Selasa": c1, "Senin": c1, "Rabu": c2, "Jumat": c2, "Kamis": c3}
    RED = "FF0000"
    BLACKF = PatternFill("solid", start_color="000000", end_color="000000")

    r = 1
    for i, line in enumerate(m.jadwal_judul[:3]):
        ws.row_dimensions[r].height = 9.2 * SCALE
        _put(ws, r, 1, line, _font("Tahoma", 7.0), border=None, c2=last_col)
        r += 1
    ws.row_dimensions[r].height = 6 * SCALE
    r += 1

    groups = [(["Selasa", "Rabu", "Kamis"], True), (["Senin", "Jumat", None], False)]
    for days, top in groups:
        lays = [_expand(p.hari[d]["baris"]) if d else [] for d in days]
        nr = max(len(x) for x in lays)
        ref = lays[0]
        rt, rl, r_rows = r, r + 1, r + 2
        ws.row_dimensions[rt].height = (8.4 if top else 12.2) * SCALE
        ws.row_dimensions[rl].height = (12.4 if top else 12.2) * SCALE
        # sumbu kiri
        _put(ws, rt, 1, "NO" if not top else None, _font("Calibri", 5.4), r2=rl if top else rt)
        _put(ws, rt, 2, "WAKTU", _font("Calibri", 5.4))
        _put(ws, rl, 2, "KELAS", _font("Calibri", 7.0 if top else 5.4))
        if not top:
            _put(ws, rl, 1, None)
        for k in range(nr):
            ws.row_dimensions[r_rows + k].height = 12.2 * SCALE
        for k, (kind, row, h) in enumerate(ref):
            rr = r_rows + k
            if kind in ("jp", "tutup_jp"):
                jam = row["jam"][h] if h < len(row["jam"]) else ""
                w = row["waktu"][h] if h < len(row["waktu"]) else ""
                _put(ws, rr, 1, jam, _font("Calibri", 5.4))
                _put(ws, rr, 2, w, _font("Calibri", 5.4))
            else:
                _put(ws, rr, 1, row.get("waktu", ""), _font("Tahoma", 5.4, True, RED), c2=2)
        # blok hari
        for di, hari in enumerate(days):
            if hari is None:
                _panel(ws, p, r, rt, rl, r_rows, nr, c3, n)
                continue
            s = blocks[hari]
            judul = p.hari[hari].get("judul") or JUDUL_HARI.get(hari, hari.upper())
            _put(ws, rt, s, judul, _font("Tahoma", 7.0 if top else 5.4), c2=s + n - 1)
            for i, kls in enumerate(p.kelas):
                _put(ws, rl, s + i, kls.nama, _font("Tahoma", 6.6 if top else 5.4), fill=_fill("#DCE6F0"), rot=90)
            lay = lays[di]
            for k, (kind, row, h) in enumerate(lay):
                rr = r_rows + k
                if kind == "info":
                    fill = _fill(row.get("warna"))
                    mt = max(n - 2, 1)
                    _put(ws, rr, s, row["label"], _font("Tahoma", 5.4 if row.get("warna") else 4.8, True,
                                                        "000000" if row.get("warna") else RED),
                         fill=fill, c2=s + mt - 1)
                    for i in range(mt, n):
                        _put(ws, rr, s + i, None, fill=fill)
                elif kind == "jeda":
                    _put(ws, rr, s, row.get("label") or None, _font("Tahoma", 5.4), c2=s + n - 1)
                elif kind in ("tutup", "tutup_jp"):
                    _put(ws, rr, s, None, fill=BLACKF, c2=s + n - 1)
                elif kind == "jp":
                    jp0 = h == 0
                    for i, kls in enumerate(p.kelas):
                        kodes = p.get_cell(hari, row["id"], kls.id)[h]
                        txt = "/".join(kodes)
                        a = by_kode.get(kodes[-1].upper()) if kodes else None
                        fnt = _kode_font(txt) if txt else "Arial-Bold"
                        _put(ws, rr, s + i, txt or None,
                             _font("Tahoma" if fnt.startswith("Tahoma") else "Arial", 5.4),
                             fill=_fill(a.warna) if a else (_fill("#FFD9D9") if txt else None))
        # pemisah hitam Selasa|Rabu
        for rr in range(rt, r_rows + nr):
            ws.cell(rr, c1 + n).fill = BLACKF
        r = r_rows + nr
    # tanda tangan
    t = m.jadwal_ttd
    ws.row_dimensions[r].height = 8 * SCALE
    r += 1
    _put(ws, r, c2, t.get("tempat_tanggal", ""), _font("Tahoma", 5.4, False), border=None, align="left", c2=c2 + n - 1)
    r += 1
    _put(ws, r, c2, t.get("jabatan", ""), _font("Tahoma", 5.4), border=None, align="left", c2=c2 + n - 1)
    r += 1
    ws.row_dimensions[r].height = 28 * SCALE
    _img(ws, m.ttd_png, f"{L(c2 + 1)}{r}", 48, 26)
    r += 1
    _put(ws, r, c2, t.get("nama", ""), _font("Tahoma", 3.8), border=None, align="left", c2=c2 + n - 1)
    r += 1
    _put(ws, r, c2, t.get("nip", ""), _font("Tahoma", 3.8), border=None, align="left", c2=c2 + n - 1)
    ws.page_setup.paperSize = 14
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{L(last_col)}{r}"
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.3
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False


def _panel(ws, p, r, rt, rl, r_rows, nr, c3, n):
    """Panel waktu Jumat di kanan-bawah (kolom ke-3)."""
    last = c3 + n - 1
    _put(ws, rt, c3, "NO", _font("Calibri", 5.4), align="left")
    _put(ws, rt, c3 + 1, "WAKTU", _font("Calibri", 5.4), c2=last)
    _put(ws, rl, c3, "KELAS", _font("Calibri", 5.4), c2=last)
    for k, (kind, row, h) in enumerate(_expand(p.hari["Jumat"]["baris"])):
        rr = r_rows + k
        if kind == "jp":
            _put(ws, rr, c3, row["jam"][h], _font("Calibri", 5.4), align="left")
            _put(ws, rr, c3 + 1, row["waktu"][h], _font("Calibri", 5.4), align="left", c2=last)
        elif kind in ("info", "jeda"):
            _put(ws, rr, c3, row.get("waktu", ""), _font("Tahoma", 5.4, True, "FF0000"), c2=last)
        else:
            _put(ws, rr, c3, row.get("label", ""), _font("Calibri", 5.4, True, "FF0000"), c2=last)


def export_xlsx(p: Proyek, path: str, halaman=("beban", "jadwal")):
    wb = Workbook()
    _sheet_beban(wb, p)
    _sheet_jadwal(wb, p)
    if "beban" not in halaman:
        wb.remove(wb["Beban Mengajar"])
    if "jadwal" not in halaman:
        wb.remove(wb["Jadwal"])
    wb.save(path)
